"""Reading an uploaded file and cutting it up.

Everything here works on bytes, so the files are built in the test rather
than kept beside it: a fixture nobody can read is a fixture nobody notices
has gone stale.
"""

import io
import zipfile

import pytest

from app.services import documents
from app.services.documents import (
    CHUNK_CHARS,
    Chunk,
    DocumentError,
    Section,
    chunk,
    clean_filename,
    extract,
    parse,
)


def _pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    count = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(count))
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Contents {5 + 2 * i} 0 R "
                "/Resources << /Font << /F1 3 0 R >> >> >>"
            ).encode()
        )
        objects.append(
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF"
    ).encode()
    return bytes(out)


def _xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Tahlilga tayyorgarlik", level=1)
    document.add_paragraph("Tahlildan oldin 8 soat ovqat yemang.")
    document.add_paragraph("Suv ichish mumkin.")
    document.add_heading("Narxlar", level=1)
    table = document.add_table(rows=3, cols=2)
    cells = [("Xizmat", "Narx"), ("UZI", "150000"), ("EKG", "80000")]
    for row, (name, price) in zip(table.rows, cells, strict=True):
        row.cells[0].text = name
        row.cells[1].text = price
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- reading -----------------------------------------------------------------


def test_plain_text_is_read_as_utf8() -> None:
    [section] = extract("eslatma.txt", "Buyrak tekshiruvi — shanba".encode())

    assert section.text == "Buyrak tekshiruvi — shanba"


def test_a_cyrillic_file_saved_by_windows_is_not_mojibake() -> None:
    """Notepad and Excel write Cyrillic as Windows-1251, and a file that
    "opens fine" there is nonsense if it is assumed to be UTF-8."""
    text = "Қабул соат тўққизда бошланади"
    russian = "Приём начинается в девять"

    [uz] = extract("a.txt", text.encode("utf-8"))
    [ru] = extract("b.txt", russian.encode("cp1251"))
    [utf16] = extract("c.txt", russian.encode("utf-16"))

    assert uz.text == text
    assert ru.text == russian
    assert utf16.text == russian


def test_markdown_headings_label_what_follows_them() -> None:
    data = b"Kirish matni\n# Narxlar\nUZI 150000\n## Tayyorgarlik\nOvqat yemang"

    sections = extract("qollanma.md", data)

    assert [(s.label, s.text) for s in sections] == [
        (None, "Kirish matni"),
        ("Narxlar", "UZI 150000"),
        ("Tayyorgarlik", "Ovqat yemang"),
    ]


@pytest.mark.parametrize("delimiter", [",", ";", "\t"])
def test_a_price_list_is_written_one_labelled_row_per_line(delimiter: str) -> None:
    """"150000" is not an answer; "Xizmat: UZI; Narx: 150000" is."""
    rows = [["Xizmat", "Narx"], ["UZI", "150000"], ["EKG", "80000"]]
    data = "\n".join(delimiter.join(row) for row in rows).encode()

    [section] = extract("narxlar.csv", data)

    assert section.table is True
    assert section.text == "Xizmat: UZI; Narx: 150000\nXizmat: EKG; Narx: 80000"


def test_a_spreadsheet_is_read_sheet_by_sheet_and_numbers_lose_their_decimal() -> None:
    data = _xlsx(
        {
            "Narxlar": [["Xizmat", "Narx"], ["UZI", 150000.0], ["", ""], ["EKG", 80000.5]],
            "Shifokorlar": [["Ism", "Kunlar"], ["Karimov", "Du-Ju"]],
        }
    )

    prices, doctors = extract("klinika.xlsx", data)

    assert prices.label == "Narxlar"
    assert prices.text == "Xizmat: UZI; Narx: 150000\nXizmat: EKG; Narx: 80000.5"
    assert (doctors.label, doctors.text) == ("Shifokorlar", "Ism: Karimov; Kunlar: Du-Ju")


def test_a_word_document_keeps_its_headings_and_reads_its_tables() -> None:
    sections = extract("qollanma.docx", _docx())

    assert [(s.label, s.table) for s in sections] == [
        ("Tahlilga tayyorgarlik", False),
        ("Narxlar", True),
    ]
    assert "Tahlildan oldin 8 soat ovqat yemang." in sections[0].text
    assert sections[1].text == "Xizmat: UZI; Narx: 150000\nXizmat: EKG; Narx: 80000"


def test_a_pdf_is_read_page_by_page() -> None:
    sections = extract("narxlar.pdf", _pdf(["UZI narxi 150000", "EKG narxi 80000"]))

    assert [(s.label, s.text) for s in sections] == [
        ("1-bet", "UZI narxi 150000"),
        ("2-bet", "EKG narxi 80000"),
    ]


# --- refusing ----------------------------------------------------------------


@pytest.mark.parametrize(
    ("filename", "data", "fragment"),
    [
        ("virus.exe", b"MZ", "formati qo'llab-quvvatlanmaydi"),
        ("eski.doc", b"\xd0\xcf", "formati qo'llab-quvvatlanmaydi"),
        ("nomsiz", b"matn", "qo'llab-quvvatlanmaydi"),
        ("bosh.txt", b"", "bo'sh"),
        ("soxta.pdf", b"bu pdf emas", "PDF emas"),
        ("buzilgan.docx", b"bu zip emas", "buzilgan"),
        ("buzilgan.xlsx", b"bu zip emas", "buzilgan"),
    ],
)
def test_a_file_that_is_not_what_it_says_is_refused_in_plain_words(
    filename: str, data: bytes, fragment: str
) -> None:
    with pytest.raises(DocumentError, match=fragment):
        extract(filename, data)


def test_a_file_over_the_size_limit_is_refused() -> None:
    with pytest.raises(DocumentError, match="juda katta"):
        extract("katta.txt", b"x" * (documents.MAX_FILE_BYTES + 1))


def test_a_small_zip_that_unpacks_to_a_flood_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * (documents.MAX_UNCOMPRESSED_BYTES + 1))

    assert len(buffer.getvalue()) < documents.MAX_FILE_BYTES
    with pytest.raises(DocumentError, match="juda katta hajmdagi"):
        extract("bomba.docx", buffer.getvalue())


def test_a_password_protected_pdf_is_refused() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(DocumentError, match="parol"):
        extract("yopiq.pdf", buffer.getvalue())


def test_a_pdf_with_no_text_says_it_may_be_a_scan() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(DocumentError, match="Skanerlangan"):
        parse("skaner.pdf", buffer.getvalue())


def test_a_parser_failure_names_the_file_not_the_parser() -> None:
    corrupt = b"%PDF-1.4\n1 0 obj\n<< /Broken"

    with pytest.raises(DocumentError) as raised:
        extract("buzuq.pdf", corrupt)

    assert "buzuq.pdf" in str(raised.value)
    assert "Traceback" not in str(raised.value)


def test_a_file_that_makes_too_many_chunks_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(documents, "MAX_CHUNKS", 3)
    text = "\n".join(f"{n}. " + "so'z " * 40 for n in range(40))

    with pytest.raises(DocumentError, match="bir necha faylga"):
        parse("uzun.txt", text.encode())


def test_a_filename_is_reduced_to_a_name() -> None:
    assert clean_filename("C:\\Users\\Ali\\Desktop\\narxlar.pdf") == "narxlar.pdf"
    assert clean_filename("../../etc/passwd") == "passwd"
    assert clean_filename("na\x00rx\x1blar.pdf") == "narxlar.pdf"
    assert clean_filename("") == "fayl"
    assert len(clean_filename("a" * 400 + ".pdf")) == 255


# --- cutting -----------------------------------------------------------------


def test_every_chunk_fits_and_no_line_is_lost() -> None:
    lines = [f"Qator {n}: " + "matn " * (n % 30) for n in range(200)]

    chunks = chunk([Section("Bet", "\n".join(lines))])

    assert all(len(c.text) <= CHUNK_CHARS for c in chunks)
    assert all(c.label == "Bet" for c in chunks)
    joined = "\n".join(c.text for c in chunks)
    assert all(line.strip() in joined for line in lines)


def test_prose_is_cut_with_an_overlap_so_a_boundary_loses_nothing() -> None:
    lines = [f"Gap raqami {n:03d} " + "a" * 60 for n in range(60)]

    chunks = chunk([Section(None, "\n".join(lines))], size=400, overlap=150)

    assert len(chunks) > 2
    for before, after in zip(chunks, chunks[1:], strict=False):
        assert before.text.splitlines()[-1] in after.text.splitlines()


def test_table_rows_are_never_split_or_overlapped() -> None:
    rows = [f"Xizmat: xizmat {n:03d}; Narx: {n * 1000}" for n in range(120)]

    chunks = chunk([Section("Narxlar", "\n".join(rows), table=True)], size=300)

    seen = [line for c in chunks for line in c.text.splitlines()]
    assert seen == rows


def test_one_enormous_line_is_cut_at_word_boundaries() -> None:
    word = "so'z"
    chunks = chunk([Section(None, " ".join([word] * 1000))], size=100, overlap=0)

    assert all(len(c.text) <= 100 for c in chunks)
    assert all(part == word for c in chunks for part in c.text.split(" "))


def test_a_single_token_longer_than_a_chunk_is_still_cut() -> None:
    chunks = chunk([Section(None, "x" * 250)], size=100, overlap=0)

    assert [len(c.text) for c in chunks] == [100, 100, 50]


def test_empty_sections_make_no_chunks() -> None:
    assert chunk([Section("Bet", ""), Section(None, "  \n  ")]) == []


def test_parse_returns_chunks_ready_to_embed() -> None:
    chunks = parse("narxlar.csv", b"Xizmat,Narx\nUZI,150000")

    assert chunks == [Chunk(None, "Xizmat: UZI; Narx: 150000")]
