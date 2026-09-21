"""Reading a file a clinic uploaded, and cutting it into pieces a question can find.

A knowledge-base row is one question and its answer, embedded by the
question. A file is not shaped like that: the answer to "UZI qancha?" is one
line of a price list, and the answer to "tahlilga qanday tayyorlanish kerak?"
is a paragraph on page three. So a file is read into sections, and the
sections are cut into chunks small enough to be found by a question and to
sit in a prompt beside the other chunks.

Two things decide whether the chunks are any good.

A row of a price list means nothing without its header: "150 000" is not an
answer, "Xizmat: UZI; Narx: 150 000" is. Spreadsheets and tables are therefore
written out one line per row, each cell labelled with its column, and packed
by whole lines so a row is never cut in two.

And prose is cut with a little overlap, so a sentence that straddles a
boundary is whole in one chunk or the other.

Nothing here talks to the database or to a model: bytes in, chunks out, and
a DocumentError -- whose message is written for the operator who uploaded the
file -- for anything that cannot be used.
"""

import csv
import io
import logging
import re
import zipfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import PurePath
from typing import Any

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 10 * 1024 * 1024
# A .docx or .xlsx is a zip, and a small zip can hold a very large file.
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_SPREADSHEET_ROWS = 20_000
MAX_CHUNKS = 1500

# Long enough to hold a paragraph or a dozen price rows, short enough that
# several fit in a prompt without crowding out the conversation.
CHUNK_CHARS = 900
CHUNK_OVERLAP = 150

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md")


class DocumentError(Exception):
    """A file that cannot be used. The message is shown to the operator."""


@dataclass(frozen=True)
class Section:
    """A run of text with one place in the file: a page, a sheet, a heading."""

    label: str | None
    text: str
    # Tables are cut between rows and never overlapped: every row already
    # carries its own column names, so nothing is lost at a boundary.
    table: bool = False


@dataclass(frozen=True)
class Chunk:
    label: str | None
    text: str


# --- reading -----------------------------------------------------------------


def clean_filename(filename: str) -> str:
    """The name to store: no directories, no control characters, 255 at most."""
    name = PurePath(filename.replace("\\", "/")).name
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return name[:255] or "fayl"


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _decode(data: bytes) -> str:
    """Text from bytes of unknown encoding.

    UTF-8 first; a UTF-16 byte-order mark next; then Windows-1251, which is
    what Excel and Notepad write when a Uzbek or Russian file is saved in
    Cyrillic and is the reason a file that "opens fine" arrives as mojibake.
    """
    candidates = ["utf-8-sig"]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        candidates.append("utf-16")
    candidates.append("cp1251")
    for encoding in candidates:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M").replace(" 00:00", "")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _table_text(rows: Sequence[Sequence[str]]) -> str:
    """A table as one labelled line per row: "Xizmat: UZI; Narx: 150000"."""
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""
    header, *body = rows
    if not body:
        return " | ".join(cell for cell in header if cell)
    lines: list[str] = []
    for row in body:
        cells: list[str] = []
        for index, value in enumerate(row):
            if not value:
                continue
            name = header[index] if index < len(header) else ""
            cells.append(f"{name}: {value}" if name else value)
        if cells:
            lines.append("; ".join(cells))
    return "\n".join(lines)


def _read_text(data: bytes, extension: str) -> list[Section]:
    text = _decode(data)
    if extension != ".md":
        return [Section(None, _clean(text))]
    # Markdown headings are the file's own table of contents: each one
    # labels what follows it.
    sections: list[Section] = []
    label: str | None = None
    body: list[str] = []
    for line in text.splitlines():
        heading = re.match(r"^#{1,6}\s+(.*\S)\s*$", line)
        if heading:
            sections.append(Section(label, _clean("\n".join(body))))
            label, body = heading.group(1)[:200], []
        else:
            body.append(line)
    sections.append(Section(label, _clean("\n".join(body))))
    return sections


def _read_csv(data: bytes) -> list[Section]:
    text = _decode(data)
    try:
        dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    if len(rows) > MAX_SPREADSHEET_ROWS:
        raise DocumentError(f"Jadval juda katta: {MAX_SPREADSHEET_ROWS} qatordan oshmasligi kerak.")
    return [Section(None, _table_text(rows), table=True)]


def _check_zip(data: bytes, extension: str) -> None:
    """Refuse a zip that is not what its name says, or that unpacks to a flood."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise DocumentError("Fayl ichida juda katta hajmdagi ma'lumot bor.")
    except zipfile.BadZipFile as error:
        raise DocumentError(
            f"Fayl «{extension}» emas yoki buzilgan. Word/Excel'da qayta saqlab yuklang."
        ) from error


def _read_xlsx(data: bytes) -> list[Section]:
    from openpyxl import load_workbook

    _check_zip(data, ".xlsx")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: list[Section] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([_cell(value) for value in row])
                if len(rows) > MAX_SPREADSHEET_ROWS:
                    raise DocumentError(
                        f"«{sheet.title}» varag'i juda katta: "
                        f"{MAX_SPREADSHEET_ROWS} qatordan oshmasligi kerak."
                    )
            sections.append(Section(sheet.title[:200], _table_text(rows), table=True))
    finally:
        workbook.close()
    return sections


def _read_docx(data: bytes) -> list[Section]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip(data, ".docx")
    document = Document(io.BytesIO(data))
    sections: list[Section] = []
    heading: str | None = None
    body: list[str] = []

    def flush() -> None:
        text = _clean("\n".join(body))
        if text:
            sections.append(Section(heading, text))
        body.clear()

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            text = paragraph.text.strip()
            if text and (style.startswith("Heading") or style == "Title"):
                flush()
                heading = text[:200]
            elif text:
                body.append(text)
        elif tag == "tbl":
            flush()
            table = Table(child, document)
            rows: list[list[str]] = []
            for row in table.rows:
                cells: list[str] = []
                for cell in row.cells:
                    value = cell.text.strip()
                    # A merged cell is reported once per column it spans.
                    if not cells or value != cells[-1]:
                        cells.append(value)
                rows.append(cells)
            sections.append(Section(heading, _table_text(rows), table=True))
    flush()
    return sections


def _read_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError("PDF parol bilan himoyalangan. Parolsiz saqlab yuklang.")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise DocumentError(f"PDF juda uzun: {MAX_PDF_PAGES} betdan oshmasligi kerak.")
    return [
        Section(f"{number}-bet", _clean(page.extract_text() or ""))
        for number, page in enumerate(reader.pages, start=1)
    ]


def extract(filename: str, data: bytes) -> list[Section]:
    """The file's text, in sections. DocumentError for anything unusable."""
    extension = PurePath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError(
            f"«{extension or filename}» formati qo'llab-quvvatlanmaydi. Yuklash mumkin: "
            + ", ".join(e.lstrip(".").upper() for e in SUPPORTED_EXTENSIONS)
            + "."
        )
    if not data:
        raise DocumentError("Fayl bo'sh.")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(
            f"Fayl juda katta ({len(data) / 1024 / 1024:.1f} MB). "
            f"Eng ko'pi {MAX_FILE_BYTES // 1024 // 1024} MB."
        )
    if extension == ".pdf" and b"%PDF" not in data[:1024]:
        raise DocumentError("Fayl PDF emas yoki buzilgan.")

    try:
        if extension == ".pdf":
            return _read_pdf(data)
        if extension == ".docx":
            return _read_docx(data)
        if extension == ".xlsx":
            return _read_xlsx(data)
        if extension == ".csv":
            return _read_csv(data)
        return _read_text(data, extension)
    except DocumentError:
        raise
    except Exception as error:
        # A parser's own message is for a developer; the operator is told
        # which file could not be read and the cause goes to the log.
        logger.exception("document_unreadable filename=%s", filename)
        raise DocumentError(
            f"«{filename}» faylini o'qib bo'lmadi. Fayl buzilgan bo'lishi mumkin."
        ) from error


# --- cutting -----------------------------------------------------------------


def _split_long(line: str, size: int) -> Iterator[str]:
    """A line longer than a chunk, cut at word boundaries."""
    buffer = ""
    for word in line.split(" "):
        while len(word) > size:
            if buffer:
                yield buffer
                buffer = ""
            yield word[:size]
            word = word[size:]
        if not buffer:
            buffer = word
        elif len(buffer) + 1 + len(word) <= size:
            buffer += " " + word
        else:
            yield buffer
            buffer = word
    if buffer:
        yield buffer


def _joined_length(lines: Sequence[str]) -> int:
    return sum(len(line) for line in lines) + max(len(lines) - 1, 0)


def _tail(lines: Sequence[str], overlap: int) -> list[str]:
    """The last whole lines, up to `overlap` characters."""
    kept: list[str] = []
    for line in reversed(lines):
        if _joined_length([line, *kept]) > overlap:
            break
        kept.insert(0, line)
    return kept


def _pack(lines: Sequence[str], size: int, overlap: int) -> list[str]:
    units = [piece for line in lines for piece in _split_long(line, size)]
    chunks: list[str] = []
    current: list[str] = []
    for unit in units:
        if current and _joined_length([*current, unit]) > size:
            chunks.append("\n".join(current))
            current = _tail(current, overlap)
            if _joined_length([*current, unit]) > size:
                current = []
        current.append(unit)
    if current:
        chunks.append("\n".join(current))
    return chunks


def chunk(
    sections: Sequence[Section], *, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP
) -> list[Chunk]:
    chunks: list[Chunk] = []
    for section in sections:
        lines = [line for line in section.text.split("\n") if line.strip()]
        for text in _pack(lines, size, 0 if section.table else overlap):
            chunks.append(Chunk(section.label, text))
    return chunks


def parse(filename: str, data: bytes) -> list[Chunk]:
    """A file as chunks, ready to embed. This is the one call the rest of the
    application makes; it is synchronous and CPU-bound, so callers run it in
    a thread."""
    chunks = chunk(extract(filename, data))
    if not chunks:
        raise DocumentError(
            "Fayldan matn topilmadi. Skanerlangan (rasm) PDF bo'lishi mumkin — "
            "matnli fayl yuklang."
        )
    if len(chunks) > MAX_CHUNKS:
        raise DocumentError(
            f"Fayl juda katta: {len(chunks)} bo'lakka bo'lindi, {MAX_CHUNKS} dan oshmasligi "
            "kerak. Uni bir necha faylga bo'lib yuklang."
        )
    return chunks
